"""
Data exporter for CSV and XLSX formats
"""
from typing import Dict, List, Optional
from datetime import datetime
from pathlib import Path
import csv
import json


class ExcelExporter:
    """Export data to Excel format"""
    
    def __init__(self, export_path: str = "./exports"):
        self.export_path = Path(export_path)
        self.export_path.mkdir(parents=True, exist_ok=True)
    
    def export_session(self, session_data: Dict, 
                      session_id: int,
                      filename: Optional[str] = None) -> Optional[str]:
        """
        Export session data to Excel
        
        Args:
            session_data: Session data dict from database
            session_id: Session ID
            filename: Custom filename (optional)
        
        Returns:
            Path to exported file or None if failed
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            print("openpyxl not installed. Skipping Excel export.")
            return None
        
        try:
            if filename is None:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"session_{session_id}_{timestamp}.xlsx"
            
            filepath = self.export_path / filename
            
            # Create workbook
            wb = openpyxl.Workbook()
            wb.remove(wb.active)  # Remove default sheet
            
            # Colors
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            session = session_data.get('session', {})
            events = session_data.get('events', [])
            analytics = session_data.get('analytics', [])
            
            # Sheet 1: Summary
            ws_summary = wb.create_sheet("Resumen")
            
            summary_data = [
                ["RESUMEN DE SESIÓN", ""],
                ["", ""],
                ["Fecha inicio", session.get('start_time', 'N/A')],
                ["Fecha fin", session.get('end_time', 'N/A')],
                ["Duración (segundos)", session.get('duration_seconds', 0)],
                ["Cámara", session.get('camera_source', 'N/A')],
                ["Notas", session.get('notes', '')],
                ["", ""],
                ["ESTADÍSTICAS GENERALES", ""],
                ["Total de eventos", len(events)],
                ["Vehículos", sum(1 for e in events if e.get('class_label') != 'Persona')],
                ["Peatones", sum(1 for e in events if e.get('class_label') == 'Persona')],
            ]
            
            # Add summary by class
            class_counts = {}
            for event in events:
                cls = event.get('class_label', 'Unknown')
                class_counts[cls] = class_counts.get(cls, 0) + 1
            
            summary_data.append(["", ""])
            summary_data.append(["POR CATEGORÍA", ""])
            for cls, count in sorted(class_counts.items()):
                summary_data.append([cls, count])
            
            # Add summary by direction
            direction_counts = {}
            for event in events:
                d = event.get('direction', 'Unknown')
                direction_counts[d] = direction_counts.get(d, 0) + 1
            
            summary_data.append(["", ""])
            summary_data.append(["POR SENTIDO", ""])
            for direction, count in sorted(direction_counts.items()):
                summary_data.append([direction, count])
            
            # Write summary sheet
            for row in summary_data:
                ws_summary.append(row)
            
            ws_summary.column_dimensions['A'].width = 25
            ws_summary.column_dimensions['B'].width = 25
            
            # Sheet 2: Events
            ws_events = wb.create_sheet("Eventos")
            headers = [
                'Timestamp',
                'Frame ID',
                'Track ID',
                'Clase',
                'Confianza',
                'Línea',
                'Sentido',
                'Posición X',
                'Posición Y'
            ]
            ws_events.append(headers)
            
            # Style header
            for cell in ws_events[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
            
            # Add events
            for event in events:
                ws_events.append([
                    event.get('timestamp', ''),
                    event.get('frame_id', ''),
                    event.get('track_id', ''),
                    event.get('class_label', ''),
                    f"{event.get('confidence', 0):.2f}",
                    event.get('line_name', ''),
                    event.get('direction', ''),
                    f"{event.get('position_x', 0):.1f}",
                    f"{event.get('position_y', 0):.1f}"
                ])
            
            # Auto-width columns
            for column in ws_events.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                ws_events.column_dimensions[column_letter].width = min(max_length + 2, 50)
            
            # Sheet 3: Analytics
            if analytics:
                ws_analytics = wb.create_sheet("Analítica")
                headers = [
                    'Inicio Intervalo',
                    'Fin Intervalo',
                    'Intervalo (min)',
                    'Total',
                    'Vehículos',
                    'Peatones',
                    'Sentido A->B',
                    'Sentido B->A'
                ]
                ws_analytics.append(headers)
                
                for cell in ws_analytics[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = border
                
                for record in analytics:
                    ws_analytics.append([
                        record.get('interval_start', ''),
                        record.get('interval_end', ''),
                        record.get('interval_minutes', ''),
                        record.get('vehicle_count', 0) + record.get('pedestrian_count', 0),
                        record.get('vehicle_count', 0),
                        record.get('pedestrian_count', 0),
                        record.get('direction_ab', 0),
                        record.get('direction_ba', 0)
                    ])
            
            # Save workbook
            wb.save(filepath)
            print(f"Excel export saved to: {filepath}")
            return str(filepath)
        
        except Exception as e:
            print(f"Error exporting to Excel: {e}")
            return None


class CSVExporter:
    """Export data to CSV format"""
    
    def __init__(self, export_path: str = "./exports"):
        self.export_path = Path(export_path)
        self.export_path.mkdir(parents=True, exist_ok=True)
    
    def export_events(self, events: List[Dict],
                     session_id: int,
                     filename: Optional[str] = None) -> Optional[str]:
        """Export events to CSV"""
        try:
            if filename is None:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"events_{session_id}_{timestamp}.csv"
            
            filepath = self.export_path / filename
            
            if not events:
                return None
            
            # Get all possible keys
            all_keys = set()
            for event in events:
                all_keys.update(event.keys())
            
            fieldnames = sorted(list(all_keys))
            
            with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
                writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(events)
            
            print(f"CSV export saved to: {filepath}")
            return str(filepath)
        
        except Exception as e:
            print(f"Error exporting to CSV: {e}")
            return None
    
    def export_summary(self, summary_data: Dict,
                      session_id: int,
                      filename: Optional[str] = None) -> Optional[str]:
        """Export summary statistics to CSV"""
        try:
            if filename is None:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"summary_{session_id}_{timestamp}.csv"
            
            filepath = self.export_path / filename
            
            with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
                writer = csv.writer(csvfile)
                
                # Write summary
                writer.writerow(['RESUMEN', ''])
                for key, value in summary_data.items():
                    if isinstance(value, dict):
                        writer.writerow([key, ''])
                        for subkey, subvalue in value.items():
                            writer.writerow([f'  {subkey}', subvalue])
                    else:
                        writer.writerow([key, value])
            
            print(f"Summary export saved to: {filepath}")
            return str(filepath)
        
        except Exception as e:
            print(f"Error exporting summary to CSV: {e}")
            return None


class DataExporter:
    """Main exporter class"""
    
    def __init__(self, export_path: str = "./exports"):
        self.export_path = export_path
        self.excel_exporter = ExcelExporter(export_path)
        self.csv_exporter = CSVExporter(export_path)
    
    def export_session_complete(self, session_data: Dict,
                                session_id: int,
                                formats: List[str] = None) -> Dict[str, str]:
        """
        Export complete session in multiple formats
        
        Args:
            session_data: Session data from database
            session_id: Session ID
            formats: List of formats ('excel', 'csv', etc.)
        
        Returns:
            Dict of format -> filepath
        """
        if formats is None:
            formats = ['excel', 'csv']
        
        results = {}
        
        if 'excel' in formats:
            filepath = self.excel_exporter.export_session(session_data, session_id)
            if filepath:
                results['excel'] = filepath
        
        if 'csv' in formats:
            events = session_data.get('events', [])
            filepath = self.csv_exporter.export_events(events, session_id)
            if filepath:
                results['csv'] = filepath
        
        return results
